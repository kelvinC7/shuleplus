"""
Modern Excel photo dashboard: uploaded, not uploaded, orphan photos,
LIST 2027 students without an image in the app.

Usage:
  python build_photo_dashboard.py --school "ST MATTHEWS TYAZO"
  python build_photo_dashboard.py --school "ST MATTHEWS TYAZO" --audit path/to/AUDIT.xlsx
"""

from __future__ import annotations

import argparse
import re
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = Path(__file__).resolve().parent

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
TITLE_FONT = Font(bold=True, size=14, color="1F4E79")
KPI_LABEL = Font(bold=True, color="404040")
ALT_ROW = PatternFill("solid", fgColor="F2F7FB")
GOOD_FILL = PatternFill("solid", fgColor="C6EFCE")
WARN_FILL = PatternFill("solid", fgColor="FFEB9C")
BAD_FILL = PatternFill("solid", fgColor="FFC7CE")


def normalize_name(name: str) -> str:
    return re.sub(r"\s+", " ", str(name).upper().strip())


def scalar(value, default=None):
    """One cell from a DataFrame row; avoids ambiguous Series truthiness."""
    if isinstance(value, pd.Series):
        if value.empty:
            return default
        value = value.iloc[0]
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return default
    return value


def audit_path_for(school: str, school_dir: Path | None = None) -> Path:
    root = school_dir or (BASE / school)
    in_folder = root / f"{school} - PHOTO UPLOAD AUDIT.xlsx"
    legacy = BASE / f"{school} - PHOTO UPLOAD AUDIT.xlsx"
    if in_folder.is_file():
        return in_folder
    if legacy.is_file():
        return legacy
    return in_folder


def load_audit(audit_file: Path | None, school: str, school_dir: Path | None = None) -> pd.DataFrame:
    path = audit_file or audit_path_for(school, school_dir)
    if not path.is_file():
        return pd.DataFrame()
    try:
        return pd.read_excel(path, sheet_name="All Rows Detail", engine="openpyxl")
    except Exception:
        return pd.read_excel(path, sheet_name=0, engine="openpyxl")


def style_sheet(ws, table_rows: int, freeze_row: int = 2) -> None:
    if ws.max_row < 1:
        return
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = ws.cell(row=freeze_row, column=1)
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{max(table_rows, ws.max_row)}"
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        max_len = 12
        for row in range(1, min(ws.max_row, 500) + 1):
            val = ws.cell(row=row, column=col).value
            if val is not None:
                max_len = max(max_len, min(len(str(val)), 48))
        ws.column_dimensions[letter].width = max_len + 2
    for row in range(2, ws.max_row + 1):
        if row % 2 == 0:
            for col in range(1, ws.max_column + 1):
                ws.cell(row=row, column=col).fill = ALT_ROW


def write_df_sheet(writer: pd.ExcelWriter, name: str, df: pd.DataFrame) -> None:
    safe = name[:31]
    df.to_excel(writer, sheet_name=safe, index=False)


def build_tyazo_list_2027_frame(school_dir: Path) -> pd.DataFrame:
    from build_list_2027_photo_status import (  # noqa: WPS433
        class_sort_key,
        enrich_image_status,
        iter_list_2027_workbooks,
        load_list_2027_students,
        load_photo_list_registry,
        load_photos_on_disk,
        match_student_to_prior_year,
    )

    photos_on_disk = load_photos_on_disk(school_dir / "photo")
    _, by_class_name, by_name = load_photo_list_registry(school_dir)
    list_2027 = load_list_2027_students(school_dir)
    rows: list[dict] = []
    for student in list_2027:
        row = match_student_to_prior_year(student, by_class_name, by_name)
        enrich_image_status([row], photos_on_disk)
        rows.append(row)
    if not rows:
        return pd.DataFrame()
    df = pd.DataFrame(rows)
    df["name_key"] = df["name"].map(normalize_name)
    return df


def load_upload_workbook_rows(school: str) -> pd.DataFrame:
    path = BASE / school / "classes with photo_id.xlsx"
    if not path.is_file():
        return pd.DataFrame()
    photo_id_min = 1000 if "RUSIZI" in school.upper() else 1800
    records: list[dict] = []
    xl = pd.ExcelFile(path, engine="openpyxl")
    for sheet in xl.sheet_names:
        raw = pd.read_excel(path, sheet_name=sheet, header=None, engine="openpyxl")
        for _, row in raw.iterrows():
            vals = [row[c] for c in raw.columns if pd.notna(row[c])]
            if len(vals) < 2:
                continue
            photo_id = None
            number = None
            name = None
            for v in vals:
                if isinstance(v, (int, float)) and float(v).is_integer():
                    iv = int(v)
                    if photo_id_min <= iv <= 4000:
                        photo_id = iv
                    elif 0 < iv < 500:
                        number = iv
                elif isinstance(v, str) and len(v.strip()) > 2:
                    upper = v.upper()
                    if upper not in {"NO", "NAMES", "PHOTO_ID"}:
                        name = v.strip()
            if name and photo_id:
                records.append(
                    {
                        "class": sheet,
                        "name": name,
                        "name_key": normalize_name(name),
                        "photo_id": photo_id,
                        "photo_file": f"2L7A{photo_id}.JPG",
                    }
                )
    if not records:
        return pd.DataFrame()
    df = pd.DataFrame(records)
    return df.drop_duplicates(subset=["name_key", "photo_id"], keep="first")


def merge_list_with_audit(list_df: pd.DataFrame, audit_df: pd.DataFrame) -> pd.DataFrame:
    if list_df.empty:
        return list_df
    out = list_df.copy()
    out["upload_status"] = ""
    out["db_matched_name"] = ""
    out["app_has_photo"] = "N"
    out["pipeline_category"] = ""

    audit_by_key: dict[tuple[str, int | None], dict] = {}
    if not audit_df.empty:
        for _, a in audit_df.iterrows():
            rec = a.to_dict()
            name = scalar(rec.get("excel_name")) or scalar(rec.get("name")) or ""
            pid = scalar(rec.get("photo_id"))
            try:
                pid_int = int(pid) if pid is not None and pid != "" else None
            except (TypeError, ValueError):
                pid_int = None
            audit_by_key[(normalize_name(str(name)), pid_int)] = rec
            if pid_int is not None:
                audit_by_key.setdefault((normalize_name(str(name)), None), rec)

    for i, row in out.iterrows():
        pid = scalar(row.get("photo_id"))
        try:
            pid_int = int(pid) if pid is not None and pid != "" else None
        except (TypeError, ValueError):
            pid_int = None
        name_key = scalar(row.get("name_key"), "")
        key = (name_key, pid_int)
        hit = audit_by_key.get(key)
        if hit is None:
            hit = audit_by_key.get((name_key, None))

        if hit is not None:
            upload_st = scalar(hit.get("upload_status"), "")
            out.at[i, "upload_status"] = upload_st or ""
            out.at[i, "db_matched_name"] = scalar(hit.get("db_matched_name"), "") or ""
            out.at[i, "app_has_photo"] = "Y" if upload_st == "Uploaded" else "N"
            if upload_st == "Uploaded":
                out.at[i, "pipeline_category"] = "Uploaded in app"
            elif scalar(hit.get("status")) == "missing_file":
                out.at[i, "pipeline_category"] = "Missing photo file"
            elif scalar(hit.get("status")) == "not_found":
                out.at[i, "pipeline_category"] = "Not in database"
            elif scalar(hit.get("ready_to_upload")) == "Y":
                out.at[i, "pipeline_category"] = "Ready to upload"
            else:
                out.at[i, "pipeline_category"] = scalar(hit.get("not_uploaded_reason")) or "Not uploaded"
        elif pid_int:
            if scalar(row.get("has_image")) == "N":
                out.at[i, "pipeline_category"] = "Missing photo file"
            else:
                out.at[i, "pipeline_category"] = "Has file — run audit on server for app status"
        else:
            out.at[i, "pipeline_category"] = "No photo_id (2027 list only)"

    return out


def orphan_photos(school_dir: Path, used_ids: set[int]) -> pd.DataFrame:
    photo_dir = school_dir / "photo"
    rows: list[dict] = []
    if not photo_dir.is_dir():
        return pd.DataFrame(columns=["photo_id", "photo_file", "status"])
    pat = re.compile(r"^2L7A(\d+)$", re.I)
    for path in sorted(photo_dir.iterdir()):
        if not path.is_file():
            continue
        m = pat.match(path.stem)
        if not m:
            continue
        pid = int(m.group(1))
        if pid not in used_ids:
            rows.append({"photo_id": pid, "photo_file": path.name, "status": "Not linked to upload list"})
    return pd.DataFrame(rows)


def kpi_row(label: str, value, note: str = "") -> dict:
    return {"Metric": label, "Value": value, "Notes": note}


def generate_dashboard(
    school: str,
    school_dir: Path | None = None,
    audit_file: Path | None = None,
) -> Path:
    school_dir = school_dir or (BASE / school)
    audit_df = load_audit(audit_file, school, school_dir)
    upload_df = load_upload_workbook_rows(school)

    used_ids: set[int] = set()
    list_df = pd.DataFrame()
    if school == "ST MATTHEWS TYAZO":
        list_df = build_tyazo_list_2027_frame(school_dir)
        list_merged = merge_list_with_audit(list_df, audit_df)
        try:
            from build_list_2027_photo_status import load_photo_list_registry  # noqa: WPS433

            photo_records, _, _ = load_photo_list_registry(school_dir)
            for rec in photo_records:
                if rec.get("photo_id"):
                    used_ids.add(int(rec["photo_id"]))
        except Exception:
            pass
    elif school in ("ST MATTHEWS RUSIZI PRIMARY", "ST MATTHEWS RUSIZI SECONDARY"):
        from build_rusizi_primary_class_photo_lists import build_rusizi_list_2027_frame  # noqa: WPS433

        list_df = build_rusizi_list_2027_frame(school_dir)
        list_merged = merge_list_with_audit(list_df, audit_df)
        if not list_df.empty and "photo_id" in list_df.columns:
            for x in list_df["photo_id"].dropna():
                try:
                    used_ids.add(int(x))
                except (TypeError, ValueError):
                    pass
    else:
        list_merged = pd.DataFrame()

    if not upload_df.empty:
        used_ids.update(int(x) for x in upload_df["photo_id"].dropna())
    if not list_df.empty and "photo_id" in list_df.columns:
        for x in list_df["photo_id"].dropna():
            try:
                used_ids.add(int(x))
            except (TypeError, ValueError):
                pass

    remain_df = orphan_photos(school_dir, used_ids)
    photo_count = len(list((school_dir / "photo").glob("*.JPG"))) if (school_dir / "photo").is_dir() else 0

    uploaded_audit = audit_df[audit_df["upload_status"] == "Uploaded"] if not audit_df.empty else pd.DataFrame()
    not_uploaded_audit = audit_df[audit_df["upload_status"] == "Not Uploaded"] if not audit_df.empty else pd.DataFrame()

    no_image_2027 = pd.DataFrame()
    by_class_summary = pd.DataFrame()
    if not list_merged.empty:
        no_image_2027 = list_merged[list_merged["app_has_photo"] != "Y"].copy()
        cols = [
            "class_2026_2027",
            "number",
            "name",
            "gender",
            "photo_id",
            "photo_file",
            "has_image",
            "pipeline_category",
            "upload_status",
            "db_matched_name",
            "match_type",
            "list_file",
        ]
        keep = [c for c in cols if c in no_image_2027.columns]
        no_image_2027 = no_image_2027[keep].sort_values(
            by=[c for c in ["class_2026_2027", "number"] if c in keep],
            kind="stable",
            na_position="last",
        )
        if "class_2026_2027" in list_merged.columns:
            by_class_summary = (
                list_merged.groupby("class_2026_2027", dropna=False)
                .agg(
                    students=("name", "count"),
                    with_photo_id=("photo_id", lambda s: s.notna().sum()),
                    in_app=("app_has_photo", lambda s: (s == "Y").sum()),
                    no_image_in_app=("app_has_photo", lambda s: (s != "Y").sum()),
                )
                .reset_index()
                .sort_values("class_2026_2027")
            )

    n_list = len(list_merged) if not list_merged.empty else 0
    n_no_app = len(no_image_2027) if not no_image_2027.empty else n_list
    n_uploaded = len(uploaded_audit)
    n_not_up = len(not_uploaded_audit)
    n_remain = len(remain_df)

    dashboard_rows = [
        kpi_row("School", school),
        kpi_row("Report generated (UTC)", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M")),
        kpi_row("", ""),
        kpi_row("JPG files in photo/", photo_count),
        kpi_row("Rows in classes with photo_id.xlsx", len(upload_df)),
        kpi_row("LIST 2027 students", n_list),
        kpi_row("Uploaded in app (from audit)", n_uploaded, "Re-run audit after upload on server"),
        kpi_row("Not uploaded (from audit)", n_not_up),
        kpi_row("LIST 2027 without image in app", n_no_app),
        kpi_row("Remaining / unlinked photos on disk", n_remain),
        kpi_row("Audit file used", str(audit_file or audit_path_for(school)), "Copy audit from server if empty"),
    ]

    out_path = school_dir / f"{school} - Photo Dashboard.xlsx"
    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        pd.DataFrame(dashboard_rows).to_excel(writer, sheet_name="Dashboard", index=False)
        if not by_class_summary.empty:
            by_class_summary.to_excel(writer, sheet_name="By Class Summary", index=False)

        up_cols = ["excel_name", "class", "photo_id", "photo_file", "db_matched_name", "match_score"]
        if not uploaded_audit.empty:
            write_df_sheet(writer, "Uploaded", uploaded_audit[[c for c in up_cols if c in uploaded_audit.columns]])
        else:
            pd.DataFrame({"info": ["No audit data — run upload-school-photos.mjs audit on server"]}).to_excel(
                writer, sheet_name="Uploaded", index=False
            )

        nu_cols = ["excel_name", "class", "photo_id", "photo_file", "not_uploaded_reason", "status", "ready_to_upload"]
        if not not_uploaded_audit.empty:
            write_df_sheet(
                writer, "Not Uploaded", not_uploaded_audit[[c for c in nu_cols if c in not_uploaded_audit.columns]]
            )
        else:
            pd.DataFrame({"info": ["No audit data"]}).to_excel(writer, sheet_name="Not Uploaded", index=False)

        if not remain_df.empty:
            remain_df.to_excel(writer, sheet_name="Remaining Photos", index=False)
        else:
            pd.DataFrame({"info": ["All disk photos linked to upload list IDs"]}).to_excel(
                writer, sheet_name="Remaining Photos", index=False
            )

        if not no_image_2027.empty:
            no_image_2027.to_excel(writer, sheet_name="LIST 2027 No Image", index=False)
        elif not list_merged.empty:
            pd.DataFrame({"info": ["All LIST 2027 rows marked with app photo"]}).to_excel(
                writer, sheet_name="LIST 2027 No Image", index=False
            )

        if not upload_df.empty:
            upload_df.to_excel(writer, sheet_name="Upload Source", index=False)

        dup_audit = (
            audit_df[audit_df["duplicate_photo_id"] == "Y"]
            if not audit_df.empty and "duplicate_photo_id" in audit_df.columns
            else pd.DataFrame()
        )
        if not dup_audit.empty:
            dup_cols = [
                c
                for c in [
                    "excel_name",
                    "class",
                    "photo_id",
                    "photo_file",
                    "other_names_same_photo_id",
                    "upload_status",
                    "db_matched_name",
                ]
                if c in dup_audit.columns
            ]
            dup_audit[dup_cols].to_excel(writer, sheet_name="Duplicate Photo ID", index=False)
        else:
            pd.DataFrame({"info": ["No duplicate photo_id in audit"]}).to_excel(
                writer, sheet_name="Duplicate Photo ID", index=False
            )

    wb = load_workbook(out_path)
    for name in wb.sheetnames:
        ws = wb[name]
        if name == "Dashboard":
            ws["A1"].font = TITLE_FONT
            for row in range(2, ws.max_row + 1):
                metric = ws.cell(row=row, column=1).value
                val = ws.cell(row=row, column=2).value
                if metric and "without image" in str(metric).lower():
                    ws.cell(row=row, column=2).fill = WARN_FILL if val else GOOD_FILL
                if metric and "Uploaded in app" in str(metric):
                    ws.cell(row=row, column=2).fill = GOOD_FILL if val and int(val) > 0 else WARN_FILL
                if metric and "Remaining" in str(metric):
                    ws.cell(row=row, column=2).fill = WARN_FILL if val and int(val) > 0 else GOOD_FILL
        style_sheet(ws, ws.max_row)
    wb.save(out_path)
    return out_path


def main() -> None:
    parser = argparse.ArgumentParser(description="Build modern photo dashboard Excel for a school")
    parser.add_argument("--school", default="ST MATTHEWS TYAZO")
    parser.add_argument("--audit", type=Path, default=None, help="Path to PHOTO UPLOAD AUDIT.xlsx")
    args = parser.parse_args()
    path = generate_dashboard(args.school, audit_file=args.audit)
    print(f"Saved: {path}")


if __name__ == "__main__":
    main()
