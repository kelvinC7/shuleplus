"""
Inventory and first-pass organization for ST MATTHEWS RUSIZI PRIMARY:
LIST 2027 (Edupoto), PRIMARY DATA WITH PHOTO (2025-2026), photo/ folder.

Usage:
  python analyze_rusizi_primary.py
"""

from __future__ import annotations

import re
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = Path(__file__).resolve().parent
SCHOOL = "ST MATTHEWS RUSIZI PRIMARY"
SCHOOL_DIR = BASE / SCHOOL
PHOTO_ID_MIN = 900
PHOTO_ID_MAX = 4000
PHOTO_PATTERN = re.compile(r"^2L7A(\d+)$", re.I)

HEADER_FILL = PatternFill("solid", fgColor="2E4057")
HEADER_FONT = Font(color="FFFFFF", bold=True)
ALT_ROW = PatternFill("solid", fgColor="F4F6F8")


def normalize_name(name: str) -> str:
    return re.sub(r"\s+", " ", str(name).upper().strip())


def class_from_list_filename(path: Path) -> str:
    return path.stem.upper().replace("_", " ").strip()


def load_photos_on_disk() -> dict[int, str]:
    photo_dir = SCHOOL_DIR / "photo"
    out: dict[int, str] = {}
    if not photo_dir.is_dir():
        return out
    for p in photo_dir.iterdir():
        if not p.is_file():
            continue
        m = PHOTO_PATTERN.match(p.stem)
        if m:
            out[int(m.group(1))] = p.name
    return out


def parse_edupoto_list_file(path: Path) -> list[dict]:
    """LIST 2027 export: lastname + firstname columns."""
    class_key = class_from_list_filename(path)
    rows: list[dict] = []
    try:
        df = pd.read_excel(path, header=None, engine="openpyxl")
    except Exception:
        df = pd.read_excel(path, header=None, engine="calamine")
    header_row = None
    for idx, row in df.iterrows():
        vals = [str(v).lower() if pd.notna(v) else "" for v in row]
        joined = " ".join(vals)
        if "lastname" in joined and "firstname" in joined:
            header_row = idx
            break
    if header_row is None:
        return rows
    header = [str(c).strip().lower() if pd.notna(c) else "" for c in df.iloc[header_row]]
    col_map = {h: i for i, h in enumerate(header) if h}

    def col(name: str) -> int | None:
        for key, i in col_map.items():
            if name in key:
                return i
        return None

    i_last = col("lastname")
    i_first = col("firstname")
    i_no = col("no") if col("no") is not None else col("student")

    for idx in range(header_row + 1, len(df)):
        row = df.iloc[idx]
        if i_last is None or i_first is None:
            break
        last = row.iloc[i_last] if i_last < len(row) else None
        first = row.iloc[i_first] if i_first < len(row) else None
        if pd.isna(last) and pd.isna(first):
            continue
        last_s = str(last).strip() if pd.notna(last) else ""
        first_s = str(first).strip() if pd.notna(first) else ""
        if not last_s and not first_s:
            continue
        full = normalize_name(f"{last_s} {first_s}".strip())
        if full in {"NAN", "NONE"} or len(full) < 3:
            continue
        number = None
        if i_no is not None:
            v = row.iloc[i_no]
            if pd.notna(v) and str(v).replace(".", "").isdigit():
                number = int(float(v))
        rows.append(
            {
                "class_2027": class_key,
                "list_file": path.name,
                "number": number,
                "name": full.title() if full.isupper() else f"{last_s} {first_s}".strip(),
                "name_key": full,
            }
        )
    return rows


def parse_primary_data_photo_file(path: Path) -> list[dict]:
    """PRIMARY DATA WITH PHOTO: title row + No, Student ID, Lastname, Firstname, photo col."""
    rows: list[dict] = []
    try:
        df = pd.read_excel(path, header=None, engine="openpyxl")
    except Exception:
        return rows
    class_hint = path.stem
    for idx, row in df.iterrows():
        vals = [row[c] for c in df.columns if pd.notna(row[c])]
        if len(vals) == 1 and isinstance(vals[0], str) and "SMATT" in vals[0].upper():
            class_hint = vals[0]
    header_row = None
    for idx, row in df.iterrows():
        texts = [str(v).lower() if pd.notna(v) else "" for v in row]
        if "lastname" in texts and "firstname" in texts:
            header_row = idx
            break
    if header_row is None:
        return rows

    for idx in range(header_row + 1, len(df)):
        row = df.iloc[idx]
        vals = [row[c] for c in df.columns if pd.notna(row[c])]
        if len(vals) < 3:
            continue
        photo_id = None
        number = None
        strings: list[str] = []
        for v in vals:
            if isinstance(v, (int, float)) and float(v).is_integer():
                iv = int(v)
                if PHOTO_ID_MIN <= iv <= PHOTO_ID_MAX and iv > 500:
                    photo_id = iv
                elif 0 < iv < 500:
                    number = iv
            elif isinstance(v, str) and v.strip():
                strings.append(v.strip())
        if len(strings) >= 2:
            name = normalize_name(f"{strings[0]} {strings[1]}")
        elif strings:
            name = normalize_name(strings[-1])
        else:
            continue
        if len(name) < 3:
            continue
        rows.append(
            {
                "source_file": path.name,
                "class_photo_list": class_hint,
                "number": number,
                "name": name,
                "name_key": name,
                "photo_id": photo_id,
                "photo_file": f"2L7A{photo_id}.JPG" if photo_id else "",
            }
        )
    return rows


def style_workbook(path: Path) -> None:
    wb = load_workbook(path)
    for ws in wb.worksheets:
        if ws.max_row < 1:
            continue
        for cell in ws[1]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
        for col in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(col)].width = 16
        for row in range(2, ws.max_row + 1):
            if row % 2 == 0:
                for col in range(1, ws.max_column + 1):
                    ws.cell(row=row, column=col).fill = ALT_ROW
    wb.save(path)


def main() -> None:
    """Legacy entry — prefer build_rusizi_primary_class_photo_lists.py for full class reports."""
    try:
        from build_rusizi_primary_class_photo_lists import main as build_class_lists

        build_class_lists()
        return
    except ImportError:
        pass

    list_dir = SCHOOL_DIR / "LIST 2027"
    photo_data_dir = SCHOOL_DIR / "PRIMARY DATA WITH PHOTO"
    photos = load_photos_on_disk()

    list_students: list[dict] = []
    if list_dir.is_dir():
        for path in sorted(list_dir.glob("*.xlsx")):
            if path.name.startswith("~$"):
                continue
            list_students.extend(parse_edupoto_list_file(path))

    prior_students: list[dict] = []
    if photo_data_dir.is_dir():
        for path in sorted(photo_data_dir.glob("*.xlsx")):
            if path.name.startswith("~$"):
                continue
            prior_students.extend(parse_primary_data_photo_file(path))

    by_name: dict[str, dict] = {}
    for rec in prior_students:
        if rec.get("photo_id"):
            by_name[rec["name_key"]] = rec

    merged: list[dict] = []
    for st in list_students:
        hit = by_name.get(st["name_key"])
        pid = hit["photo_id"] if hit else None
        merged.append(
            {
                **st,
                "photo_id": pid,
                "photo_file": f"2L7A{pid}.JPG" if pid else "",
                "has_prior_photo_id": "Y" if pid else "N",
                "has_jpg": "Y" if pid and pid in photos else "N",
                "match_type": "exact_name" if hit else "none",
                "prior_class_file": hit["source_file"] if hit else "",
            }
        )

    used_ids = {r["photo_id"] for r in prior_students if r.get("photo_id")}
    orphan = [
        {"photo_id": pid, "photo_file": fname}
        for pid, fname in sorted(photos.items())
        if pid not in used_ids
    ]

    list_files = sorted(list_dir.glob("*.xlsx")) if list_dir.is_dir() else []
    photo_files = sorted(photo_data_dir.glob("*.xlsx")) if photo_data_dir.is_dir() else []

    inventory = pd.DataFrame(
        [
            {"Item": "School", "Count / Value": SCHOOL},
            {"Item": "Generated (UTC)", "Count / Value": datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M")},
            {"Item": "LIST 2027 class files", "Count / Value": len(list_files)},
            {"Item": "LIST 2027 students", "Count / Value": len(list_students)},
            {"Item": "PRIMARY DATA WITH PHOTO files", "Count / Value": len(photo_files)},
            {"Item": "Prior list rows with photo_id", "Count / Value": len([r for r in prior_students if r.get("photo_id")])},
            {"Item": "JPG in photo/", "Count / Value": len(photos)},
            {"Item": "2027 students matched photo_id (exact name)", "Count / Value": sum(1 for m in merged if m["photo_id"])},
            {"Item": "2027 with photo_id + JPG on disk", "Count / Value": sum(1 for m in merged if m["has_jpg"] == "Y")},
            {"Item": "2027 still no image", "Count / Value": sum(1 for m in merged if m["has_jpg"] != "Y")},
            {"Item": "Unlinked JPGs (not in prior lists)", "Count / Value": len(orphan)},
        ]
    )

    file_map = []
    for p in list_files:
        file_map.append({"Area": "LIST 2027", "File": p.name, "Class key": class_from_list_filename(p)})
    for p in photo_files:
        file_map.append({"Area": "PRIMARY DATA WITH PHOTO", "File": p.name, "Class key": p.stem})

    by_class = pd.DataFrame(merged)
    if not by_class.empty:
        class_summary = (
            by_class.groupby("class_2027")
            .agg(
                students=("name_key", "count"),
                with_photo_id=("photo_id", lambda s: s.notna().sum()),
                with_jpg=("has_jpg", lambda s: (s == "Y").sum()),
                no_image=("has_jpg", lambda s: (s != "Y").sum()),
            )
            .reset_index()
        )
    else:
        class_summary = pd.DataFrame()

    no_image_df = by_class[by_class["has_jpg"] != "Y"] if not by_class.empty else pd.DataFrame()

    out = SCHOOL_DIR / f"{SCHOOL} - Organization Report.xlsx"
    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        inventory.to_excel(writer, sheet_name="Dashboard", index=False)
        pd.DataFrame(file_map).to_excel(writer, sheet_name="File Map", index=False)
        if not class_summary.empty:
            class_summary.to_excel(writer, sheet_name="By Class Summary", index=False)
        if not by_class.empty:
            by_class.to_excel(writer, sheet_name="LIST 2027 With Photo ID", index=False)
        if not no_image_df.empty:
            no_image_df.to_excel(writer, sheet_name="LIST 2027 No Image", index=False)
        pd.DataFrame(prior_students).to_excel(writer, sheet_name="Prior Photo Lists", index=False)
        pd.DataFrame(orphan).to_excel(writer, sheet_name="Remaining Photos", index=False)

    style_workbook(out)

    upload_ready = [m for m in merged if m.get("photo_id") and m.get("has_jpg") == "Y"]
    if upload_ready:
        by_class: dict[str, list] = {}
        for row in upload_ready:
            by_class.setdefault(row["class_2027"], []).append(row)
        upload_path = SCHOOL_DIR / "classes with photo_id.xlsx"
        with pd.ExcelWriter(upload_path, engine="openpyxl") as writer:
            for class_name in sorted(by_class.keys()):
                sheet = class_name[:31]
                matrix = [["NO", "NAMES", "photo_id"]]
                for r in sorted(by_class[class_name], key=lambda x: (x.get("number") is None, x.get("number") or 999)):
                    matrix.append([r.get("number") or "", r["name"], r["photo_id"]])
                pd.DataFrame(matrix).to_excel(writer, sheet_name=sheet, index=False, header=False)
        print(f"  Upload workbook: {upload_path.name} ({len(upload_ready)} rows)")

    print(SCHOOL)
    print(f"  LIST 2027 students: {len(list_students)}")
    print(f"  Prior rows with photo_id: {len([r for r in prior_students if r.get('photo_id')])}")
    print(f"  JPG on disk: {len(photos)}")
    print(f"  2027 matched photo_id: {sum(1 for m in merged if m['photo_id'])}")
    print(f"  2027 with JPG: {sum(1 for m in merged if m['has_jpg'] == 'Y')}")
    print(f"  Saved: {out.name}")


if __name__ == "__main__":
    main()
