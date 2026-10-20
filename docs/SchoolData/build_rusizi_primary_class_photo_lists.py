"""
ST MATTHEWS RUSIZI PRIMARY: compare LIST 2027 with PRIMARY DATA WITH PHOTO + photo/.
Writes per-class workbooks (Has Photo / No Photo) and summary reports.

Usage:
  python build_rusizi_primary_class_photo_lists.py
"""

from __future__ import annotations

import json
import re
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = Path(__file__).resolve().parent
SCHOOL = "ST MATTHEWS RUSIZI PRIMARY"
SCHOOL_DIR = BASE / SCHOOL
PHOTO_ID_MIN = 1000
PHOTO_ID_MAX = 4000
PHOTO_PATTERN = re.compile(r"^2L7A(\d+)$", re.I)
SKIP_SCAN_DIRS = {
    "LIST 2027",
    "LIST 2027 By Class Photo Status",
    "LIST 2027 By Class Has Photo",
    "LIST 2027 with photo_id",
}

GENERATED_REPORT_SUFFIXES = (
    " - 2027 By Class Has Photo and No Photo.xlsx",
    " - LIST 2027 With Photo and App Status.xlsx",
    " - Photo Dashboard.xlsx",
    " - Students and Staff With Photos.xlsx",
    " - Organization Report.xlsx",
    " - PHOTO UPLOAD AUDIT.xlsx",
    " - PHOTO UPLOAD REPORT.xlsx",
)

HEADER_FILL = PatternFill("solid", fgColor="1B4332")
HEADER_FONT = Font(color="FFFFFF", bold=True)
GOOD_FILL = PatternFill("solid", fgColor="D8F3DC")
WARN_FILL = PatternFill("solid", fgColor="FFF3B0")
ALT_ROW = PatternFill("solid", fgColor="F8F9FA")


def normalize_name(name: str) -> str:
    return re.sub(r"\s+", " ", str(name).upper().strip())


def class_from_list_filename(path: Path) -> str:
    return path.stem.upper().replace("_", " ").strip()


def load_photos_on_disk() -> dict[int, str]:
    photo_dir = SCHOOL_DIR / "photo"
    out: dict[int, str] = {}
    if not photo_dir.is_dir():
        return out
    for p in photo_dir.iterdir():
        if not p.is_file():
            continue
        m = PHOTO_PATTERN.match(p.stem)
        if m:
            out[int(m.group(1))] = p.name
    return out


def parse_edupoto_list_file(path: Path) -> list[dict]:
    class_key = class_from_list_filename(path)
    rows: list[dict] = []
    try:
        df = pd.read_excel(path, header=None, engine="openpyxl")
    except Exception:
        df = pd.read_excel(path, header=None, engine="calamine")

    header_row = None
    for idx, row in df.iterrows():
        joined = " ".join(str(v).lower() for v in row if pd.notna(v))
        if "lastname" in joined and "firstname" in joined:
            header_row = idx
            break
    if header_row is None:
        return rows

    header = [str(c).strip().lower() if pd.notna(c) else "" for c in df.iloc[header_row]]

    def find_col(*needles: str) -> int | None:
        for i, h in enumerate(header):
            if all(n in h for n in needles):
                return i
        return None

    i_no = find_col("no")
    i_last = find_col("lastname")
    i_first = find_col("firstname")
    i_gender = find_col("gender") or find_col("sex")
    if i_last is None or i_first is None:
        return rows

    for idx in range(header_row + 1, len(df)):
        row = df.iloc[idx]
        last = row.iloc[i_last] if i_last < len(row) else None
        first = row.iloc[i_first] if i_first < len(row) else None
        if pd.isna(last) and pd.isna(first):
            continue
        last_s = str(last).strip() if pd.notna(last) else ""
        first_s = str(first).strip() if pd.notna(first) else ""
        if not last_s and not first_s:
            continue
        display = f"{last_s} {first_s}".strip()
        name_key = normalize_name(display)
        if len(name_key) < 3:
            continue
        number = None
        if i_no is not None and i_no < len(row):
            v = row.iloc[i_no]
            if pd.notna(v):
                try:
                    number = int(float(v))
                except (TypeError, ValueError):
                    pass
        gender = ""
        if i_gender is not None and i_gender < len(row):
            g = row.iloc[i_gender]
            if pd.notna(g):
                gender = str(g).strip()
        rows.append(
            {
                "class_2027": class_key,
                "list_file": path.name,
                "number": number,
                "name": display,
                "name_key": name_key,
                "gender": gender,
            }
        )
    return rows


def parse_csv_list_file(path: Path) -> list[dict]:
    """ShulePlus / Edupoto CSV export (First Name, Last Name, Class, …)."""
    class_key = path.stem.upper().replace("_", " ").strip()
    try:
        df = pd.read_csv(path, encoding="utf-8-sig")
    except UnicodeDecodeError:
        df = pd.read_csv(path, encoding="latin-1")

    col_map = {str(c).strip().lower(): c for c in df.columns}

    def col(*names: str) -> str | None:
        for n in names:
            if n in col_map:
                return col_map[n]
        return None

    c_no = col("no", "no.")
    c_first = col("first name", "firstname", "first")
    c_last = col("last name", "lastname", "last", "surname")
    c_gender = col("gender", "sex")
    c_class = col("class", "classname", "grade")
    if not c_first and not c_last:
        return []

    rows: list[dict] = []
    for _, row in df.iterrows():
        first_s = str(row[c_first]).strip() if c_first and pd.notna(row.get(c_first)) else ""
        last_s = str(row[c_last]).strip() if c_last and pd.notna(row.get(c_last)) else ""
        if not first_s and not last_s:
            continue
        display = normalize_name(f"{first_s} {last_s}")
        if len(display) < 3:
            continue
        class_label = class_key
        if c_class and pd.notna(row.get(c_class)):
            class_label = str(row[c_class]).strip().upper() or class_key
        number = None
        if c_no and pd.notna(row.get(c_no)):
            try:
                number = int(float(row[c_no]))
            except (TypeError, ValueError):
                pass
        gender = ""
        if c_gender and pd.notna(row.get(c_gender)):
            gender = str(row[c_gender]).strip()
        rows.append(
            {
                "class_2027": class_label,
                "list_file": path.name,
                "number": number,
                "name": display,
                "name_key": display,
                "gender": gender,
            }
        )
    return rows


def iter_list_2027_files(list_dir: Path) -> list[Path]:
    """Class lists in LIST 2027/ (and nested LIST 2027/LIST 2027/ after Edupoto export)."""
    if not list_dir.is_dir():
        return []
    paths: list[Path] = []
    for folder in (list_dir, list_dir / "LIST 2027"):
        if not folder.is_dir():
            continue
        for path in sorted(folder.iterdir()):
            if not path.is_file() or path.name.startswith("~$"):
                continue
            if path.suffix.lower() in {".xlsx", ".csv"}:
                paths.append(path)
    seen: set[Path] = set()
    unique: list[Path] = []
    for path in paths:
        resolved = path.resolve()
        if resolved in seen:
            continue
        seen.add(resolved)
        unique.append(path)
    return unique


def load_list_2027_students(list_dir: Path) -> list[dict]:
    students: list[dict] = []
    for path in iter_list_2027_files(list_dir):
        suffix = path.suffix.lower()
        if suffix == ".xlsx":
            students.extend(parse_edupoto_list_file(path))
        elif suffix == ".csv":
            students.extend(parse_csv_list_file(path))
    return students


def name_lookup_keys(name_key: str) -> list[str]:
    """Exact key plus given/family flip (Edupoto vs ShulePlus CSV order)."""
    key = normalize_name(name_key)
    keys = [key]
    parts = key.split()
    if len(parts) >= 2:
        keys.append(normalize_name(f"{parts[-1]} {' '.join(parts[:-1])}"))
        if len(parts) == 2:
            keys.append(normalize_name(f"{parts[1]} {parts[0]}"))
    out: list[str] = []
    seen: set[str] = set()
    for k in keys:
        if k and k not in seen:
            seen.add(k)
            out.append(k)
    return out


def photo_id_from_row(row: pd.Series, header: list[str], i_no: int | None, i_last: int | None, i_first: int | None) -> int | None:
    for i, h in enumerate(header):
        if not h:
            continue
        if "photo" in h or "phot_id" in h:
            if i < len(row):
                v = row.iloc[i]
                if isinstance(v, (int, float)) and pd.notna(v) and float(v).is_integer():
                    iv = int(v)
                    if PHOTO_ID_MIN <= iv <= PHOTO_ID_MAX:
                        return iv

    skip: set[int] = set()
    for idx in (i_no, i_last, i_first):
        if idx is not None:
            skip.add(idx)
    for i, h in enumerate(header):
        if h and "student" in h and "id" in h:
            skip.add(i)

    candidates: list[tuple[int, int]] = []
    for i in range(len(row)):
        if i in skip:
            continue
        v = row.iloc[i]
        if isinstance(v, (int, float)) and pd.notna(v) and float(v).is_integer():
            iv = int(v)
            if PHOTO_ID_MIN <= iv <= PHOTO_ID_MAX:
                candidates.append((i, iv))
    if not candidates:
        return None
    return candidates[-1][1]


def parse_primary_data_photo_file(path: Path) -> list[dict]:
    rows: list[dict] = []
    try:
        df = pd.read_excel(path, header=None, engine="openpyxl")
    except Exception:
        return rows

    class_hint = path.stem
    header_row = None
    for idx, row in df.iterrows():
        vals = [row[c] for c in df.columns if pd.notna(row[c])]
        if len(vals) == 1 and isinstance(vals[0], str) and "SMATT" in vals[0].upper():
            class_hint = vals[0]
        texts = [str(v).lower() if pd.notna(v) else "" for v in row]
        if "lastname" in texts and "firstname" in texts:
            header_row = idx
            break
    if header_row is None:
        return rows

    header = [str(c).strip().lower() if pd.notna(c) else "" for c in df.iloc[header_row]]

    def col_idx(label: str) -> int | None:
        for i, h in enumerate(header):
            if label in h:
                return i
        return None

    i_no = col_idx("no")
    i_last = col_idx("lastname")
    i_first = col_idx("firstname")
    if i_last is None or i_first is None:
        return rows

    for idx in range(header_row + 1, len(df)):
        row = df.iloc[idx]
        last = row.iloc[i_last] if i_last < len(row) else None
        first = row.iloc[i_first] if i_first < len(row) else None
        if pd.isna(last) and pd.isna(first):
            continue
        last_s = str(last).strip() if pd.notna(last) else ""
        first_s = str(first).strip() if pd.notna(first) else ""
        display = f"{last_s} {first_s}".strip()
        name_key = normalize_name(display)
        if len(name_key) < 3:
            continue

        number = None
        if i_no is not None and i_no < len(row) and pd.notna(row.iloc[i_no]):
            try:
                number = int(float(row.iloc[i_no]))
            except (TypeError, ValueError):
                pass

        photo_id = photo_id_from_row(row, header, i_no, i_last, i_first)

        rows.append(
            {
                "source_file": path.name,
                "class_photo_list": class_hint,
                "number": number,
                "name": display,
                "name_key": name_key,
                "photo_id": photo_id,
                "photo_file": f"2L7A{photo_id}.JPG" if photo_id else "",
            }
        )
    return rows


def parse_classes_with_photo_id_workbook(path: Path) -> list[dict]:
    rows: list[dict] = []
    try:
        xl = pd.ExcelFile(path, engine="openpyxl")
    except Exception:
        return rows
    for sheet in xl.sheet_names:
        if sheet.lower() in {"summary", "by class summary"}:
            continue
        df = pd.read_excel(path, sheet_name=sheet, header=None, engine="openpyxl")
        for _, row in df.iterrows():
            vals = [row[c] for c in df.columns if pd.notna(row[c])]
            if len(vals) < 2:
                continue
            photo_id = None
            number = None
            names: list[str] = []
            for v in vals:
                if isinstance(v, (int, float)) and float(v).is_integer():
                    iv = int(v)
                    if PHOTO_ID_MIN <= iv <= PHOTO_ID_MAX:
                        photo_id = iv
                    elif 0 < iv < 500:
                        number = iv
                elif isinstance(v, str):
                    t = v.strip()
                    if t.upper() in {"NO", "NAMES", "PHOTO_ID", "PHOT_ID"}:
                        continue
                    if len(t) > 2:
                        names.append(t)
            if not names or not photo_id:
                continue
            name = max(names, key=len)
            name_key = normalize_name(name)
            rows.append(
                {
                    "source_file": f"{path.name} [{sheet}]",
                    "class_photo_list": sheet,
                    "number": number,
                    "name": name,
                    "name_key": name_key,
                    "photo_id": photo_id,
                    "photo_file": f"2L7A{photo_id}.JPG",
                }
            )
    return rows


def parse_pdf_marks_json(path: Path) -> list[dict]:
    """LIST WITH PHOTOID/pdf with all classes - marks photo_ids.json (handwritten IDs from marks PDF)."""
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return []
    rows: list[dict] = []
    for block in data.get("classes", []):
        class_label = block.get("class_label", "")
        for st in block.get("students", []):
            pid = st.get("photo_id")
            if pid is None:
                continue
            last = str(st.get("lastname") or "").strip()
            first = str(st.get("firstname") or "").strip()
            display = f"{last} {first}".strip()
            name_key = normalize_name(display)
            if len(name_key) < 3:
                continue
            rows.append(
                {
                    "source_file": f"{path.name} (page {block.get('pdf_page')})",
                    "class_photo_list": class_label,
                    "number": st.get("no"),
                    "name": display,
                    "name_key": name_key,
                    "photo_id": int(pid),
                    "photo_file": f"2L7A{int(pid)}.JPG",
                }
            )
    return rows


def load_pdf_marks_photo_records(school_dir: Path) -> list[dict]:
    photoid_dir = school_dir / "LIST WITH PHOTOID"
    if not photoid_dir.is_dir():
        return []
    for name in (
        "pdf with all classes - marks photo_ids.json",
        "pdf with all classes - marks photo_ids.JSON",
    ):
        path = photoid_dir / name
        if path.is_file():
            return parse_pdf_marks_json(path)
    return []


def iter_photo_excel_files(school_dir: Path) -> list[Path]:
    paths: list[Path] = []
    for path in sorted(school_dir.rglob("*.xlsx")):
        if path.name.startswith("~$"):
            continue
        if any(part in SKIP_SCAN_DIRS for part in path.parts):
            continue
        if path.parent.name == "LIST 2027":
            continue
        if "TEACHERS" in path.name.upper():
            continue
        if any(path.name.endswith(suffix) for suffix in GENERATED_REPORT_SUFFIXES):
            continue
        paths.append(path)
    return paths


def load_all_photo_id_records(school_dir: Path) -> tuple[list[dict], pd.DataFrame]:
    records: list[dict] = []
    audit_rows: list[dict] = []
    for path in iter_photo_excel_files(school_dir):
        parsed: list[dict] = []
        rel = path.relative_to(school_dir).as_posix()
        if path.name.lower() == "classes with photo_id.xlsx":
            parsed = parse_classes_with_photo_id_workbook(path)
        else:
            parsed = parse_primary_data_photo_file(path)
        audit_rows.append(
            {
                "file": rel,
                "rows_with_photo_id": sum(1 for r in parsed if r.get("photo_id")),
                "student_rows": len(parsed),
            }
        )
        records.extend(parsed)
    pdf_rows = load_pdf_marks_photo_records(school_dir)
    if pdf_rows:
        audit_rows.append(
            {
                "file": "LIST WITH PHOTOID/pdf with all classes - marks photo_ids.json",
                "rows_with_photo_id": len(pdf_rows),
                "student_rows": len(pdf_rows),
            }
        )
        records.extend(pdf_rows)
    return records, pd.DataFrame(audit_rows)


def class_stream(class_label: str) -> str:
    u = class_label.upper()
    m = re.search(r"(N[123]|P[1-6])\s*(EAST|WEST|SOUTH|NORTH|A|B|CONTACT)?", u)
    if not m:
        return u
    grade = m.group(1)
    stream = (m.group(2) or "").strip()
    return f"{grade} {stream}".strip()


def build_photo_registry(prior_students: list[dict]) -> tuple[dict[str, list[dict]], dict[str, dict]]:
    by_name: dict[str, list[dict]] = {}
    for rec in prior_students:
        if not rec.get("photo_id"):
            continue
        by_name.setdefault(rec["name_key"], []).append(rec)
    best: dict[str, dict] = {}
    return by_name, best


def resolve_best_by_name(by_name: dict[str, list[dict]], photos: dict[int, str]) -> dict[str, dict]:
    best: dict[str, dict] = {}
    for key, recs in by_name.items():
        best[key] = sorted(
            recs,
            key=lambda r: (
                1 if r.get("photo_id") in photos else 0,
                r.get("photo_id") or 0,
            ),
            reverse=True,
        )[0]
    return best


def match_student(
    student: dict,
    by_name_all: dict[str, list[dict]],
    by_name_best: dict[str, dict],
    photos: dict[int, str],
) -> dict:
    base = {**student, "photo_id": None, "photo_file": "", "has_jpg": "N", "match_type": "none", "prior_source": ""}
    hit = None
    matched_key = None
    for key in name_lookup_keys(student["name_key"]):
        hit = by_name_best.get(key)
        if hit:
            matched_key = key
            break
    if hit:
        pid = hit["photo_id"]
        base.update(
            {
                "photo_id": pid,
                "photo_file": f"2L7A{pid}.JPG",
                "has_jpg": "Y" if pid in photos else "N",
                "match_type": "exact_same_name",
                "prior_source": hit["source_file"],
                "prior_class": hit.get("class_photo_list", ""),
            }
        )
        if matched_key and len(by_name_all.get(matched_key, [])) > 1:
            base["match_type"] = "exact_name_duplicate_in_prior"
        elif matched_key != student["name_key"]:
            base["match_type"] = "exact_name_order_variant"
        return base

    return base


def style_worksheet(ws, freeze: str = "A2") -> None:
    if ws.max_row < 1:
        return
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    ws.freeze_panes = freeze
    if ws.max_row > 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    for col in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col)].width = 18
    for row in range(2, ws.max_row + 1):
        if row % 2 == 0:
            for col in range(1, ws.max_column + 1):
                ws.cell(row=row, column=col).fill = ALT_ROW


def load_audit_map(
    school_dir: Path | None = None,
    school_name: str | None = None,
) -> tuple[dict[tuple[str, int], dict], dict[int, dict]]:
    root = school_dir or SCHOOL_DIR
    name = school_name or SCHOOL
    path = root / f"{name} - PHOTO UPLOAD AUDIT.xlsx"
    if not path.is_file():
        path = BASE / f"{name} - PHOTO UPLOAD AUDIT.xlsx"
    if not path.is_file():
        return {}, {}
    try:
        audit = pd.read_excel(path, sheet_name="All Rows Detail", engine="openpyxl")
    except Exception:
        return {}, {}
    out: dict[tuple[str, int], dict] = {}
    by_photo_id: dict[int, dict] = {}
    for _, r in audit.iterrows():
        try:
            pid = int(r["photo_id"])
        except (TypeError, ValueError, KeyError):
            continue
        excel_name = r.get("excel_name", r.get("name", ""))
        key = (normalize_name(str(excel_name)), pid)
        rec = {
            "upload_status": str(r.get("upload_status") or ""),
            "db_matched_name": str(r.get("db_matched_name") or ""),
            "app_status": str(r.get("status") or ""),
            "not_uploaded_reason": str(r.get("not_uploaded_reason") or ""),
            "has_existing_photo": str(r.get("has_existing_photo") or ""),
            "match_score": str(r.get("match_score") or ""),
        }
        out[key] = rec
        if pid not in by_photo_id or rec.get("upload_status") == "Uploaded":
            by_photo_id[pid] = rec
    return out, by_photo_id


def apply_audit_columns(
    all_df: pd.DataFrame,
    audit_map: dict[tuple[str, int], dict],
    audit_by_photo_id: dict[int, dict] | None = None,
) -> pd.DataFrame:
    df = all_df.copy()
    by_pid = audit_by_photo_id or {}
    for col in (
        "upload_status",
        "db_matched_name",
        "app_status",
        "not_uploaded_reason",
        "has_existing_photo",
        "match_score",
    ):
        df[col] = ""
    for idx, r in df.iterrows():
        if pd.isna(r.get("photo_id")):
            continue
        pid = int(r["photo_id"])
        key = (r["name_key"], pid)
        hit = audit_map.get(key) or by_pid.get(pid)
        if not hit:
            continue
        for col, val in hit.items():
            df.at[idx, col] = val
    df["app_has_photo"] = ""
    for idx, r in df.iterrows():
        if pd.isna(r.get("photo_id")):
            continue
        pid = int(r["photo_id"])
        key = (r["name_key"], pid)
        hit = audit_map.get(key) or by_pid.get(pid)
        if not hit:
            continue
        if hit["upload_status"] == "Uploaded" or hit["has_existing_photo"] == "Y":
            df.at[idx, "app_has_photo"] = "Y"
        else:
            df.at[idx, "app_has_photo"] = "N"
    return df


def write_list_2027_with_photo_id(all_df: pd.DataFrame) -> Path:
    """One workbook per LIST 2027 file: NO, NAMES, PHOTO_ID, HAS_JPG for every 2027 student."""
    out_root = SCHOOL_DIR / "LIST 2027 with photo_id"
    out_root.mkdir(parents=True, exist_ok=True)
    for list_file in sorted(all_df["list_file"].unique()):
        part = all_df[all_df["list_file"] == list_file].sort_values(
            by="number", kind="stable", na_position="last"
        )
        class_title = part.iloc[0]["class_2027"]
        matrix: list[list] = [[class_title], ["NO", "NAMES", "GENDER", "PHOTO_ID", "HAS_JPG"]]
        for _, r in part.iterrows():
            pid = r.get("photo_id")
            matrix.append(
                [
                    r.get("number") or "",
                    r["name"],
                    r.get("gender") or "",
                    int(pid) if pd.notna(pid) else "",
                    r.get("has_jpg") or "N",
                ]
            )
        pd.DataFrame(matrix).to_excel(out_root / list_file, index=False, header=False)
    return out_root


def write_list_2027_has_image_app_workbook(
    all_df: pd.DataFrame,
    school_dir: Path | None = None,
    school_name: str | None = None,
) -> Path:
    """Master: LIST 2027 students with JPG on disk + upload audit columns when present."""
    root = school_dir or SCHOOL_DIR
    name = school_name or SCHOOL
    out_path = root / f"{name} - LIST 2027 With Photo and App Status.xlsx"
    has_img = all_df[all_df["has_jpg"] == "Y"].copy()
    detail_cols = [
        "class_2027",
        "number",
        "name",
        "gender",
        "photo_id",
        "photo_file",
        "list_file",
        "upload_status",
        "app_status",
        "db_matched_name",
        "has_existing_photo",
        "app_has_photo",
        "not_uploaded_reason",
        "match_score",
        "prior_source",
    ]
    detail_cols = [c for c in detail_cols if c in all_df.columns]

    uploaded = has_img[has_img["upload_status"] == "Uploaded"] if "upload_status" in has_img.columns else pd.DataFrame()
    not_in_db = has_img[has_img["app_status"] == "not_found"] if "app_status" in has_img.columns else pd.DataFrame()
    pending_app = (
        has_img[(has_img["upload_status"] != "Uploaded") & (has_img["app_status"] == "matched")]
        if "app_status" in has_img.columns and "upload_status" in has_img.columns
        else pd.DataFrame()
    )

    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        pd.DataFrame(
            [
                {"Metric": "School", "Value": name},
                {"Metric": "Generated (UTC)", "Value": datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M")},
                {"Metric": "LIST 2027 students total", "Value": len(all_df)},
                {"Metric": "LIST 2027 with photo_id", "Value": int(all_df["photo_id"].notna().sum())},
                {"Metric": "LIST 2027 with JPG (has image)", "Value": len(has_img)},
                {"Metric": "Uploaded in app (from audit)", "Value": len(uploaded)},
                {"Metric": "Has JPG — not in database", "Value": len(not_in_db)},
                {"Metric": "Has JPG — in DB, photo pending", "Value": len(pending_app)},
            ]
        ).to_excel(writer, sheet_name="Summary", index=False)
        has_img[detail_cols].sort_values(
            by=["class_2027", "number"], kind="stable", na_position="last"
        ).to_excel(writer, sheet_name="All Has Image", index=False)
        if not uploaded.empty:
            uploaded[detail_cols].to_excel(writer, sheet_name="Uploaded In App", index=False)
        if not not_in_db.empty:
            not_in_db[detail_cols].to_excel(writer, sheet_name="Import Priority", index=False)
        if not pending_app.empty:
            pending_app[detail_cols].to_excel(writer, sheet_name="In DB Photo Pending", index=False)
        all_df[detail_cols + (["match_type"] if "match_type" in all_df.columns else [])].sort_values(
            by=["class_2027", "number"], kind="stable", na_position="last"
        ).to_excel(writer, sheet_name="All LIST 2027", index=False)

    wb = load_workbook(out_path)
    for ws in wb.worksheets:
        style_worksheet(ws)
    wb.save(out_path)
    return out_path


def photo_compare_status(row: pd.Series) -> str:
    if row.get("has_jpg") == "Y":
        return "Has photo (ID + JPG)"
    if pd.notna(row.get("photo_id")):
        return "Has photo_id — JPG missing"
    return "No photo_id"


def write_list_2027_class_by_class_compare(
    all_df: pd.DataFrame,
    school_dir: Path | None = None,
    school_name: str | None = None,
    prior_students: list[dict] | None = None,
    photos: dict[int, str] | None = None,
) -> Path:
    """LIST 2027 vs confirmed photo_id sources, one sheet per class + summary."""
    root = school_dir or SCHOOL_DIR
    name = school_name or SCHOOL
    out_path = root / f"{name} - LIST 2027 Class by Class Photo Compare.xlsx"

    df = all_df.copy()
    df["photo_status"] = df.apply(photo_compare_status, axis=1)

    detail_cols = [
        "number",
        "name",
        "gender",
        "photo_id",
        "photo_file",
        "photo_status",
        "has_jpg",
        "match_type",
        "prior_source",
        "list_file",
    ]
    detail_cols = [c for c in detail_cols if c in df.columns]

    class_summary = (
        df.groupby("class_2027", sort=True)
        .agg(
            list_file=("list_file", "first"),
            students=("name", "count"),
            has_photo_id=("photo_id", lambda s: s.notna().sum()),
            has_jpg=("has_jpg", lambda s: (s == "Y").sum()),
            no_photo_id=("photo_id", lambda s: s.isna().sum()),
            jpg_missing=("photo_status", lambda s: (s == "Has photo_id — JPG missing").sum()),
        )
        .reset_index()
        .rename(columns={"class_2027": "Class (LIST 2027)"})
    )

    list_keys: set[str] = set()
    for _, r in df.iterrows():
        for k in name_lookup_keys(str(r["name_key"] if "name_key" in r else r["name"])):
            list_keys.add(k)

    orphan_rows: list[dict] = []
    if prior_students and photos is not None:
        seen_pid_name: set[tuple[int, str]] = set()
        for rec in prior_students:
            pid = rec.get("photo_id")
            if not pid:
                continue
            key = rec.get("name_key") or normalize_name(rec.get("name", ""))
            matched = any(k in list_keys for k in name_lookup_keys(key))
            if matched:
                continue
            dedupe = (int(pid), key)
            if dedupe in seen_pid_name:
                continue
            seen_pid_name.add(dedupe)
            orphan_rows.append(
                {
                    "prior_class": rec.get("class_photo_list", ""),
                    "number": rec.get("number"),
                    "name": rec.get("name"),
                    "photo_id": int(pid),
                    "photo_file": f"2L7A{int(pid)}.JPG",
                    "has_jpg": "Y" if int(pid) in photos else "N",
                    "prior_source": rec.get("source_file", ""),
                    "on_list_2027": "N",
                }
            )

    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        pd.DataFrame(
            [
                {"Metric": "School", "Value": name},
                {"Metric": "Generated (UTC)", "Value": datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M")},
                {"Metric": "LIST 2027 classes", "Value": df["class_2027"].nunique()},
                {"Metric": "LIST 2027 students", "Value": len(df)},
                {"Metric": "Has photo (ID + JPG)", "Value": int((df["photo_status"] == "Has photo (ID + JPG)").sum())},
                {"Metric": "Has photo_id — JPG missing", "Value": int((df["photo_status"] == "Has photo_id — JPG missing").sum())},
                {"Metric": "No photo_id", "Value": int((df["photo_status"] == "No photo_id").sum())},
                {
                    "Metric": "Photo_id in sources, not on LIST 2027",
                    "Value": len(orphan_rows),
                },
            ]
        ).to_excel(writer, sheet_name="Summary", index=False)
        class_summary.to_excel(writer, sheet_name="By Class Summary", index=False)

        used_sheets: set[str] = set()
        for class_name in sorted(df["class_2027"].unique()):
            part = df[df["class_2027"] == class_name].sort_values(
                by="number", kind="stable", na_position="last"
            )
            sheet = re.sub(r'[\\/*?:\[\]]', "-", str(class_name))[:31] or "Class"
            if sheet in used_sheets:
                suffix = 2
                while f"{sheet[:28]}_{suffix}" in used_sheets:
                    suffix += 1
                sheet = f"{sheet[:28]}_{suffix}"
            used_sheets.add(sheet)
            title = f"{class_name} ({part.iloc[0]['list_file']})" if not part.empty else class_name
            matrix: list[list] = [[title], detail_cols]
            for _, r in part.iterrows():
                row_vals: list = []
                for c in detail_cols:
                    v = r.get(c, "")
                    if c == "photo_id" and pd.notna(v):
                        v = int(v)
                    row_vals.append(v)
                matrix.append(row_vals)
            pd.DataFrame(matrix).to_excel(writer, sheet_name=sheet, index=False, header=False)

        if orphan_rows:
            pd.DataFrame(orphan_rows).sort_values(["prior_class", "name"]).to_excel(
                writer, sheet_name="Photo ID Not On LIST 2027", index=False
            )

    wb = load_workbook(out_path)
    for ws in wb.worksheets:
        style_worksheet(ws)
    wb.save(out_path)
    return out_path


def write_import_priority_upload_workbook(not_in_db: pd.DataFrame) -> Path | None:
    if not_in_db.empty:
        return None
    out_path = SCHOOL_DIR / "classes with photo_id - import priority.xlsx"
    class_order = sorted(not_in_db["class_2027"].unique())
    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        for class_name in class_order:
            part = not_in_db[not_in_db["class_2027"] == class_name]
            sheet = class_name[:31]
            matrix = [["NO", "NAMES", "photo_id"]]
            for _, r in part.sort_values("number", na_position="last").iterrows():
                matrix.append([r.get("number") or "", r["name"], int(r["photo_id"])])
            pd.DataFrame(matrix).to_excel(writer, sheet_name=sheet, index=False, header=False)
    return out_path


def compute_list_2027_status(
    school_dir: Path | None = None,
    school_name: str | None = None,
) -> tuple[pd.DataFrame, pd.DataFrame, list[dict], dict[int, str]]:
    root = school_dir or SCHOOL_DIR
    name = school_name or root.name
    list_dir = root / "LIST 2027"
    photo_dir = root / "photo"
    photos: dict[int, str] = {}
    if photo_dir.is_dir():
        for p in photo_dir.iterdir():
            if not p.is_file():
                continue
            m = PHOTO_PATTERN.match(p.stem)
            if m:
                photos[int(m.group(1))] = p.name
    list_students = load_list_2027_students(list_dir)
    prior_students, excel_audit = load_all_photo_id_records(root)
    by_name_all, _ = build_photo_registry(prior_students)
    by_name_best = resolve_best_by_name(by_name_all, photos)
    matched_rows = [match_student(st, by_name_all, by_name_best, photos) for st in list_students]
    all_df = pd.DataFrame(matched_rows)
    audit_map, audit_by_photo_id = load_audit_map(root, name)
    if audit_map or audit_by_photo_id:
        all_df = apply_audit_columns(all_df, audit_map, audit_by_photo_id)
    return all_df, excel_audit, prior_students, photos


def build_rusizi_list_2027_frame(school_dir: Path | None = None) -> pd.DataFrame:
    """Used by build_photo_dashboard.py for Rusizi Primary / Secondary."""
    root = school_dir or SCHOOL_DIR
    all_df, _, _, _ = compute_list_2027_status(root, root.name)
    all_df["class_2026_2027"] = all_df["class_2027"]
    all_df["has_image"] = all_df["has_jpg"]
    if "photo_id" in all_df.columns:
        all_df["photo_file"] = all_df.apply(
            lambda r: f"2L7A{int(r['photo_id'])}.JPG" if pd.notna(r.get("photo_id")) else "",
            axis=1,
        )
    return all_df


def write_class_workbook(out_path: Path, class_name: str, has_photo: pd.DataFrame, no_photo: pd.DataFrame) -> None:
    summary = pd.DataFrame(
        [
            {"Metric": "Class", "Value": class_name},
            {"Metric": "Students total", "Value": len(has_photo) + len(no_photo)},
            {"Metric": "Has photo (ID + JPG file)", "Value": len(has_photo)},
            {"Metric": "No photo (see No Photo sheet)", "Value": len(no_photo)},
        ]
    )
    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Summary", index=False)
        if not has_photo.empty:
            has_photo.to_excel(writer, sheet_name="Has Photo", index=False)
        else:
            pd.DataFrame({"info": ["None"]}).to_excel(writer, sheet_name="Has Photo", index=False)
        if not no_photo.empty:
            no_photo.to_excel(writer, sheet_name="No Photo", index=False)
        else:
            pd.DataFrame({"info": ["None"]}).to_excel(writer, sheet_name="No Photo", index=False)

    wb = load_workbook(out_path)
    for name in wb.sheetnames:
        style_worksheet(wb[name])
    wb.save(out_path)


def main() -> None:
    photos = load_photos_on_disk()
    jpg_file_count = len(list((SCHOOL_DIR / "photo").glob("*.JPG"))) if (SCHOOL_DIR / "photo").is_dir() else 0

    all_df, excel_audit, prior_students, photos = compute_list_2027_status()
    has_any_id = all_df[all_df["photo_id"].notna()].copy()
    no_id = all_df[all_df["photo_id"].isna()].copy()

    out_cols_has = [
        "number",
        "name",
        "photo_id",
        "photo_file",
        "has_jpg",
        "match_type",
        "prior_source",
        "list_file",
    ]
    out_cols_no = ["number", "name", "list_file", "class_2027"]

    per_class_dir = SCHOOL_DIR / "LIST 2027 By Class Has Photo"
    per_class_dir.mkdir(parents=True, exist_ok=True)

    class_order = sorted(all_df["class_2027"].unique())
    for class_name in class_order:
        part = all_df[all_df["class_2027"] == class_name]
        has_part = part[part["has_jpg"] == "Y"][out_cols_has].sort_values(
            by="number", kind="stable", na_position="last"
        )
        no_part = part[part["has_jpg"] != "Y"].copy()
        if not no_part.empty:
            no_part["reason"] = no_part.apply(
                lambda r: "No photo_id in any school Excel"
                if pd.isna(r.get("photo_id"))
                else "photo_id set but JPG missing in photo/",
                axis=1,
            )
            no_part = no_part[["number", "name", "list_file", "class_2027", "reason"]]
        else:
            no_part = pd.DataFrame(columns=["number", "name", "list_file", "class_2027", "reason"])
        no_part = no_part.sort_values(by="number", kind="stable", na_position="last")
        safe = re.sub(r'[<>:"/\\|?*]', "-", class_name)[:80]
        write_class_workbook(per_class_dir / f"{safe}.xlsx", class_name, has_part, no_part)

    used_ids = {int(r["photo_id"]) for r in prior_students if r.get("photo_id")}
    remain = [
        {"photo_id": pid, "photo_file": fname}
        for pid, fname in sorted(photos.items())
        if pid not in used_ids
    ]

    by_class_summary = (
        all_df.groupby("class_2027")
        .agg(
            students=("name", "count"),
            has_photo_id=("photo_id", lambda s: s.notna().sum()),
            has_jpg=("has_jpg", lambda s: (s == "Y").sum()),
            no_photo_id=("photo_id", lambda s: s.isna().sum()),
        )
        .reset_index()
    )

    master_path = SCHOOL_DIR / f"{SCHOOL} - 2027 By Class Has Photo and No Photo.xlsx"
    with pd.ExcelWriter(master_path, engine="openpyxl") as writer:
        pd.DataFrame(
            [
                {"Metric": "School", "Value": SCHOOL},
                {"Metric": "Generated (UTC)", "Value": datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M")},
                {"Metric": "LIST 2027 students", "Value": len(all_df)},
                {"Metric": "With photo_id (prior lists)", "Value": len(has_any_id)},
                {"Metric": "With JPG on disk", "Value": int((all_df["has_jpg"] == "Y").sum())},
                {"Metric": "No photo_id match", "Value": len(no_id)},
                {"Metric": "Photo_id rows (all Excel sources)", "Value": len([r for r in prior_students if r.get("photo_id")])},
                {"Metric": "Excel files scanned for photo_id", "Value": len(excel_audit)},
                {"Metric": "JPG files in photo/ (count)", "Value": jpg_file_count},
                {"Metric": "Unique 2L7A IDs on disk", "Value": len(photos)},
                {"Metric": "Unlinked JPGs", "Value": len(remain)},
            ]
        ).to_excel(writer, sheet_name="Summary", index=False)
        if not excel_audit.empty:
            excel_audit.to_excel(writer, sheet_name="Photo Excel Audit", index=False)
        by_class_summary.to_excel(writer, sheet_name="By Class Summary", index=False)
        all_df[all_df["has_jpg"] == "Y"][out_cols_has + ["class_2027"]].to_excel(
            writer, sheet_name="All Has Photo", index=False
        )
        no_all = all_df[all_df["has_jpg"] != "Y"].copy()
        no_all["reason"] = no_all.apply(
            lambda r: "No photo_id in any school Excel"
            if pd.isna(r.get("photo_id"))
            else "photo_id set but JPG missing",
            axis=1,
        )
        no_all[["number", "name", "class_2027", "list_file", "reason"]].to_excel(
            writer, sheet_name="All No Photo", index=False
        )
        has_any_id[has_any_id["has_jpg"] == "N"][out_cols_has + ["class_2027"]].to_excel(
            writer, sheet_name="Has ID Missing JPG", index=False
        )
        pd.DataFrame(remain).to_excel(writer, sheet_name="Remaining Photos", index=False)

    wb = load_workbook(master_path)
    for ws in wb.worksheets:
        style_worksheet(ws)
    wb.save(master_path)

    upload_path = SCHOOL_DIR / "classes with photo_id.xlsx"
    upload_rows = has_any_id[has_any_id["has_jpg"] == "Y"]
    upload_note = ""
    if not upload_rows.empty:
        try:
            with pd.ExcelWriter(upload_path, engine="openpyxl") as writer:
                for class_name in class_order:
                    part = upload_rows[upload_rows["class_2027"] == class_name]
                    if part.empty:
                        continue
                    sheet = class_name[:31]
                    matrix = [["NO", "NAMES", "photo_id"]]
                    for _, r in part.sort_values("number", na_position="last").iterrows():
                        matrix.append([r.get("number") or "", r["name"], int(r["photo_id"])])
                    pd.DataFrame(matrix).to_excel(writer, sheet_name=sheet, index=False, header=False)
        except OSError as exc:
            upload_note = f" (skipped: {exc})"

    list_photo_dir = write_list_2027_with_photo_id(all_df)
    app_status_path = write_list_2027_has_image_app_workbook(all_df)
    compare_path = write_list_2027_class_by_class_compare(
        all_df,
        prior_students=prior_students,
        photos=photos,
    )
    not_in_db = (
        all_df[(all_df["has_jpg"] == "Y") & (all_df["app_status"] == "not_found")]
        if "app_status" in all_df.columns
        else pd.DataFrame()
    )
    import_wb = write_import_priority_upload_workbook(not_in_db)

    print(SCHOOL)
    print(f"  LIST 2027 students: {len(all_df)}")
    print(f"  Matched photo_id: {len(has_any_id)}")
    print(f"  With JPG: {(all_df['has_jpg'] == 'Y').sum()}")
    print(f"  No photo_id: {len(no_id)}")
    print(f"  Per-class folder: {per_class_dir}")
    print(f"  Master workbook: {master_path.name}")
    print(f"  Upload source: {upload_path.name} ({len(upload_rows)} rows){upload_note}")
    print(f"  LIST 2027 with photo_id/: {list_photo_dir}")
    print(f"  App status workbook: {app_status_path.name}")
    print(f"  Class-by-class compare: {compare_path.name}")
    if import_wb:
        print(f"  Import priority upload list: {import_wb.name} ({len(not_in_db)} rows)")

    try:
        from build_photo_dashboard import generate_dashboard  # noqa: WPS433

        dash = generate_dashboard(SCHOOL, SCHOOL_DIR)
        print(f"  Photo dashboard: {dash.name}")
    except Exception as exc:
        print(f"  Photo dashboard skipped: {exc}")


if __name__ == "__main__":
    main()
